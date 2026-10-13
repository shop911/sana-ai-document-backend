"""Sana AI Document backend. No paid AI, text PDFs only, local processing."""
import io, os, re
from pathlib import Path
from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
import pdfplumber
from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from xml.sax.saxutils import escape
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

app=Flask(__name__)
app.config['MAX_CONTENT_LENGTH']=12*1024*1024
CORS(
    app,
    resources={r"/api/*": {
        "origins": [
            "http://localhost:5057",
            "http://127.0.0.1:5057"
        ]
    }}
)
ALLOWED={'.pdf','.docx','.xlsx'}

def upload(field='file'):
    f=request.files.get(field)
    if not f or not f.filename or Path(f.filename).suffix.lower() not in ALLOWED:
        raise ValueError('Upload a PDF, DOCX or XLSX file.')
    return f,Path(f.filename).suffix.lower()

def pdf_text(data):
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        return '\n'.join((page.extract_text(layout=False) or '') for page in pdf.pages)

def get_text(data, ext):
    if ext=='.pdf': return pdf_text(data)
    if ext=='.docx':
        doc=Document(io.BytesIO(data)); return '\n'.join([p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(max_row=1500, max_col=50, values_only=True):
                if any(v is not None for v in row):
                    lines.append(' | '.join(str(v) if v is not None else '' for v in row))
        return '\n'.join(lines)
    finally:
        wb.close()

def extract_parts(text):
    out=[]; seen=set()
    for line in text.splitlines():
        line=' '.join(line.split())
        if len(line)<8:continue
        # Conservative extraction: require an alphanumeric part identifier containing digits.
        tokens=line.split(); candidates=[t.strip(',:;()[]') for t in tokens if re.fullmatch(r'(?=[A-Za-z0-9./_-]*\d)[A-Za-z0-9][A-Za-z0-9./_-]{3,34}',t.strip(',:;()[]'))]
        if not candidates:continue
        part=candidates[0]
        if part in seen:continue
        seen.add(part)
        desc=line[line.find(part)+len(part):].strip(' -:|') or 'REVIEW DESCRIPTION'
        out.append({'part_no':part,'mfg_part_no':part,'description':desc[:40], 'full_description':desc,'tag':'WSS Handling','review':True})
    return out[:500]

def xlsx_bytes(headers, rows):
    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Sana Document AI'
    sheet.append(headers)
    for row in rows:
        sheet.append([str(value) if value is not None else '' for value in row])
    for cell in sheet[1]:
        cell.fill = PatternFill(fill_type='solid', fgColor='202A53')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for i, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = min(65, max(22, len(str(header)) + 3))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def pdf_excel_rows(data):
    """Extract actual PDF table cells when available; fall back to readable lines."""
    rows = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            page_tables = page.extract_tables() or []
            found = False
            for table in page_tables:
                if not table:
                    continue
                valid_rows = []
                for row in table:
                    if row and any(str(v or '').strip() for v in row):
                        valid_rows.append([str(v or '').strip()[:3000] for v in row])
                if valid_rows:
                    found = True
                    rows.append([f'Page {page_number} - Table'])
                    rows.extend(valid_rows)
                    rows.append([])
            if not found:
                page_text = page.extract_text(layout=False) or ''
                for line in page_text.splitlines():
                    if line.strip():
                        rows.append([line.strip()[:3000]])
    return rows[:4000]

def as_download(data,filename,mime):return send_file(io.BytesIO(data),as_attachment=True,download_name=filename,mimetype=mime)

@app.get('/api/health')
def health():return jsonify(ok=True,service='Sana Document AI V3')

@app.post('/api/extract')
def extract():
    try:
        f,ext=upload();txt=get_text(f.read(),ext)
        if not txt.strip():return jsonify(error='No selectable text found. Scanned PDF OCR is not available in V3.'),422
        return jsonify(rows=extract_parts(txt),preview=txt[:2200],warning='Automatic guesses: verify every part number and description against the source before export.')
    except Exception as e:return jsonify(error=str(e)),400

@app.post('/api/export-parts')
def export_parts():
    try:
        data=request.get_json(force=True); rows=data.get('rows',[])
        if len(rows)>500:return jsonify(error='Maximum 500 rows'),400
        headers=['Part No','Mfg Part No','Description (max 40)','Full Description','Area','Review Required']
        values=[[str(r.get('part_no',''))[:80],str(r.get('mfg_part_no',''))[:80],str(r.get('description',''))[:40],str(r.get('full_description',''))[:2000],str(r.get('tag','WSS Handling'))[:80], 'Yes' if r.get('review',True) else 'No'] for r in rows]
        return as_download(xlsx_bytes(headers,values),'Sana_AI_Parts_Review.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    except Exception as e:return jsonify(error=str(e)),400

@app.post('/api/convert')
def convert():
    try:
        f,ext=upload(); data=f.read(); target=request.form.get('target','xlsx').lower()
        text=get_text(data,ext)
        if not text.strip():return jsonify(error='No selectable text found. OCR for scanned PDFs is not included.'),422
        if target=='xlsx':
            if ext == '.pdf':
                rows = pdf_excel_rows(data)
            else:
                rows = [[line[:3000]] for line in text.splitlines() if line.strip()][:4000]
            if not rows:
                return jsonify(error='No text or tables found. Scanned PDFs need OCR.'),422
            width = max(len(row) for row in rows)
            headers = ['Column ' + str(i) for i in range(1, width + 1)]
            return as_download(xlsx_bytes(headers, rows),'Sana_Converted.xlsx','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        if target=='docx':
            doc=Document();doc.add_heading('Sana AI - Extracted Text',0)
            for line in text.splitlines()[:4000]:doc.add_paragraph(line)
            out=io.BytesIO();doc.save(out)
            return as_download(out.getvalue(),'Sana_Converted.docx','application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        if target=='pdf':
            out=io.BytesIO();story=[]; styles=getSampleStyleSheet()
            for line in text.splitlines()[:4000]:
                story.append(Paragraph(escape(line[:1500]) or ' ',styles['Normal']));story.append(Spacer(1,4))
            SimpleDocTemplate(out).build(story)
            return as_download(out.getvalue(),'Sana_Converted.pdf','application/pdf')
        return jsonify(error='Choose xlsx, docx or pdf'),400
    except Exception as e:return jsonify(error=str(e)),400

if __name__=='__main__':app.run(host='127.0.0.1',port=int(os.getenv('PORT','5056')))
